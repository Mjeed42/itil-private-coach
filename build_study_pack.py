#!/usr/bin/env python3
"""Build text-forward knowledge files for the mobile ITIL study agent."""

from pathlib import Path
import html
import json
import re
from docx import Document

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "itil-study-agent" / "knowledge"
SITE = ROOT / "docs"


def docx_text(path: Path) -> str:
    document = Document(path)
    blocks = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def srt_text(path: Path) -> str:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    lines = []
    previous = None
    for line in raw.splitlines():
        line = line.strip()
        if not line or line.isdigit() or "-->" in line:
            continue
        line = html.unescape(re.sub(r"<[^>]+>", "", line))
        if line != previous:
            lines.append(line)
            previous = line
    return " ".join(lines)


def natural_key(path: Path):
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", path.stem)]


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    transcript = [
        "# ITIL 4 Course Transcript",
        "",
        "> Source: caption files in this repository. Generated for retrieval; wording may contain caption errors.",
    ]
    for path in sorted((ROOT / "Subtitles").glob("*.srt"), key=natural_key):
        transcript.extend(["", f"## {path.stem}", "", srt_text(path)])
    (OUT / "course-transcript.md").write_text("\n".join(transcript) + "\n", encoding="utf-8")

    exams = [
        "# Practice Exams",
        "",
        "> Keep answers hidden until the learner commits to an answer. Preserve the source wording.",
    ]
    for filename in (
        "ITIL 4 Foundation (Practice Exam 1).docx",
        "ITIL 4 Foundation (Practice Exam 2).docx",
    ):
        exams.extend(["", f"## {filename}", "", docx_text(ROOT / filename)])
    (OUT / "practice-exams.md").write_text("\n".join(exams) + "\n", encoding="utf-8")

    metadata = docx_text(ROOT / "ITIL 4 Foundation.docx")
    (OUT / "course-metadata.md").write_text(
        "# Course Metadata\n\n" + metadata + "\n", encoding="utf-8"
    )

    voice_rules = """PRIVATE ITIL 4 FOUNDATION VOICE COACH PACK

HOW TO ACTIVATE

When this file is attached as a ChatGPT Project source, the learner will say:
"Start my ITIL 4 diagnostic using the Voice Coach pack."

Follow the REQUIRED TUTOR BEHAVIOR below for the entire project chat.

REQUIRED VOICE TUTOR BEHAVIOR

- Act as a rigorous, encouraging ITIL 4 Foundation exam coach and thinking partner.
- First confirm the booked exam is ITIL 4 Foundation, not ITIL Foundation (Version 5). Never mix versions.
- Ask for the exam date, available daily study time, preferred language, and prior IT service-management experience.
- Run a 12-question diagnostic covering key concepts, four dimensions, service value system, guiding principles, service value chain, continual improvement, and practices.
- In voice sessions, ask exactly one question at a time, speak all four options clearly as A through D, then stop and wait for the learner.
- Never reveal an answer before the learner commits. Afterward use: verdict; correct answer; why; why the tempting alternative is wrong; one memory hook.
- Keep most spoken turns below 60 seconds. Avoid spoken tables. If the learner says "repeat", restate more slowly and simply.
- Use plain English by default. If Arabic or bilingual teaching is requested, retain the official English ITIL term beside its Arabic explanation.
- Teach with retrieval practice: short explanation, workplace example, confusing contrast, learner answer, correction, later retest.
- Track errors by concept. Revisit missed concepts later in the session and again after 1, 3, and 7 days.
- Treat current official PeopleCert material as higher authority than this 2019 course. Identify legacy wording such as "change control" and provide the current equivalent when applicable.
- For full mocks: 40 questions, 60-minute target, closed-book simulation, feedback only at the end.
- Do not call a learner exam-ready until they score at least 80% on two fresh timed mocks and no major area is below 70%.
- End each session with a concise spoken summary and a written ITIL PROGRESS SNAPSHOT containing date, minutes, topics, score, strengths, weak areas, corrected misconceptions, 1/3/7-day reviews, next session, and mock history.
- Use the course transcript as teaching context. Use the practice-exam answer keys internally, but keep answers hidden until the learner responds.
- Never invent a definition, current exam policy, syllabus weighting, source, or answer key. State uncertainty and verify current policy when needed.

CURRICULUM CHECKLIST

- Key concepts: service management, value and value co-creation, stakeholders, products, services, offerings, relationships, utility, warranty, outputs, outcomes, costs, and risks.
- Four dimensions and PESTLE factors.
- Service value system, governance, opportunity, demand, value, practices, and continual improvement.
- Seven guiding principles and their interactions.
- Six service value-chain activities, purposes, inputs/outputs, and use in value streams.
- Continual improvement model and register.
- Purposes and key terms of syllabus management practices, with scenario-level mastery of the high-priority practices.
- Exam technique, mixed retrieval, weak-area repair, and timed mocks.

SOURCE PRIORITY

1. Current official PeopleCert source supplied or verified during the chat.
2. Course source material included below.
3. Generated explanations and questions, clearly distinguished from source questions.

BEGIN COURSE SOURCES
"""
    combined = "\n\n".join([
        voice_rules,
        metadata,
        "COURSE TRANSCRIPT\n\n" + "\n".join(transcript[3:]),
        "PRACTICE EXAMS AND ANSWER KEYS\n\n" + "\n".join(exams[3:]),
    ])
    (OUT.parent / "ITIL-ChatGPT-Voice-Coach.txt").write_text(combined + "\n", encoding="utf-8")

    questions = []
    for exam_number, filename in enumerate((
        "ITIL 4 Foundation (Practice Exam 1).docx",
        "ITIL 4 Foundation (Practice Exam 2).docx",
    ), start=1):
        document = Document(ROOT / filename)
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        starts = [i for i, text in enumerate(paragraphs) if re.fullmatch(r"Question \d+:", text)]
        answers = {
            int(row.cells[0].text.strip()): row.cells[1].text.strip().upper()
            for table in document.tables
            for row in table.rows
            if row.cells and row.cells[0].text.strip().isdigit()
        }
        for offset, start in enumerate(starts):
            number = int(re.search(r"\d+", paragraphs[start]).group())
            end = starts[offset + 1] if offset + 1 < len(starts) else len(paragraphs)
            block = [x for x in paragraphs[start + 1:end] if x.lower() != "answer key"]
            if len(block) < 5:
                raise ValueError(f"Could not parse {filename}, question {number}")
            prompt, options = " ".join(block[:-4]), block[-4:]
            key = answers[number]
            questions.append({
                "id": f"e{exam_number}q{number}",
                "exam": exam_number,
                "number": number,
                "prompt": prompt,
                "options": options,
                "answer": ord(key) - ord("A"),
                "topic": classify_topic(prompt + " " + " ".join(options)),
                "legacy": "change control" in (prompt + " " + " ".join(options)).lower(),
            })
    SITE.mkdir(parents=True, exist_ok=True)
    payload = "window.ITIL_QUESTIONS = " + json.dumps(questions, ensure_ascii=False, indent=2) + ";\n"
    (SITE / "questions.js").write_text(payload, encoding="utf-8")

    print(f"Built knowledge pack, Voice Coach file, and {len(questions)} site questions")


def classify_topic(text: str) -> str:
    value = text.lower()
    rules = (
        ("principles", ("guiding principle", "focus on value", "start where", "iteratively", "collaborate", "holistically", "simple and practical", "optimize and automate")),
        ("dimensions", ("four dimension", "organizations and people", "information and technology", "partners and suppliers", "value streams and processes", "pestle")),
        ("continual", ("continual improvement", "improvement register", "where are we", "vision")),
        ("value-chain", ("value chain", "plan activity", "engage activity", "obtain/build", "deliver and support", "design and transition")),
        ("practices", ("practice", "incident", "problem", "service desk", "service request", "service level", "supplier", "information security", "deployment", "release", "configuration", "asset", "monitoring", "event", "availability", "capacity", "continuity")),
        ("svs", ("service value system", "governance", "opportunity", "demand")),
        ("concepts", ("service", "utility", "warranty", "outcome", "output", "value", "risk", "cost", "customer", "sponsor", "user", "service relationship")),
    )
    for topic, keywords in rules:
        if any(keyword in value for keyword in keywords):
            return topic
    return "concepts"


if __name__ == "__main__":
    main()
